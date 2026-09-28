import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

import os
_src = "/tmp/cirurgioes.json"
CIRURGIOES = [r["name"] for r in json.load(open(_src))] if os.path.exists(_src) else []
STATUS = ["AGENDADO","AGUARDANDO","COM PENDENCIA","DESISTIU","MENSAGEM ENVIADA","NAO RESPONDE","PROBLEMA COM TELEFONE","SUSPENSA REAGENDAR","REALIZADO"]
CONVENIOS = ["SUS","PARTICULAR"]
ANESTESIAS = ["RAQUIANESTESIA","GERAL","SEDAÇÃO","LOCAL","PERIDURAL","BLOQUEIO PERIFÉRICO"]
PRIORIDADES = ["ELETIVA","URGÊNCIA","PRIORIDADE","EMERGÊNCIA"]
ESPECIALIDADES = ["ANESTESIOLOGISTA","CIRURGIA GERAL","ORTOPEDIA","VASCULAR","GINECOLOGIA","OTORRINOLARINGOLOGIA","UROLOGIA"]
SIMNAO = ["SIM","NAO"]

# (coluna, obrigatoria_negocio, exemplo, lista_validacao)
COLS = [
 ("nomePaciente", True,  "MARIA APARECIDA DE SOUZA", None),
 ("cpf",          "ou",  "12345678901", None),
 ("cns",          "ou",  "", None),
 ("nascimento",   False, "1970-05-12", None),
 ("telefone1",    False, "(15) 99999-0000", None),
 ("telefone2",    False, "", None),
 ("municipio",    False, "SÃO PAULO", None),
 ("cirurgiao",    True,  (CIRURGIOES[0] if CIRURGIOES else ""), "CIRURGIOES"),
 ("especialidade",False, "CIRURGIA GERAL", "ESPECIALIDADES"),
 ("procedimento", True,  "HERNIOPLASTIA RECIDIVANTE", None),
 ("anestesia",    False, "GERAL", "ANESTESIAS"),
 ("convenio",     False, "SUS", "CONVENIOS"),
 ("prioridade",   False, "ELETIVA", "PRIORIDADES"),
 ("sala",         False, "", None),
 ("dataAtendimento", False, "2026-06-20", None),
 ("dataAutorizacao", False, "", None),
 ("dataAgendado",    False, "", None),
 ("horario",      False, "08:30", None),
 ("aih",          False, "NAO", "SIMNAO"),
 ("autorizada",   False, "NAO", "SIMNAO"),
 ("apa",          False, "NAO", "SIMNAO"),
 ("opme",         False, "NAO", "SIMNAO"),
 ("status",       True,  "AGUARDANDO", "STATUS"),
 ("observacoes",  False, "", None),
 ("unidade",      False, "", None),
 ("duracao",      False, "60", None),
]
LISTAS = {"STATUS":STATUS,"CONVENIOS":CONVENIOS,"ANESTESIAS":ANESTESIAS,"PRIORIDADES":PRIORIDADES,"ESPECIALIDADES":ESPECIALIDADES,"SIMNAO":SIMNAO,"CIRURGIOES":CIRURGIOES}

F="Calibri"
RED=PatternFill("solid",fgColor="C00000"); BLUE=PatternFill("solid",fgColor="1F4E78"); AMBER=PatternFill("solid",fgColor="BF8F00")
HF=Font(name=F,bold=True,color="FFFFFF",size=10); EX=Font(name=F,italic=True,color="808080",size=10)
thin=Side(style="thin",color="D9D9D9"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=Workbook()
ws=wb.active; ws.title="Dados"; ws.freeze_panes="A3"
for ci,(h,ob,ex,vk) in enumerate(COLS,1):
    c=ws.cell(1,ci,h); c.font=HF; c.alignment=Alignment("center","center",wrap_text=True); c.border=BORDER
    c.fill = RED if ob is True else (AMBER if ob=="ou" else BLUE)
    e=ws.cell(2,ci,ex); e.font=EX; e.border=BORDER
    ws.column_dimensions[get_column_letter(ci)].width=max(13,min(26,len(h)+4))
    if h in ("cpf","cns","telefone1","telefone2","horario","duracao","nascimento","dataAtendimento","dataAutorizacao","dataAgendado"):
        for r in range(2,1002): ws.cell(r,ci).number_format="@"
ws.row_dimensions[1].height=30

wv=wb.create_sheet("Valores válidos")
ordem=[("STATUS","status"),("ESPECIALIDADES","especialidade"),("ANESTESIAS","anestesia"),("CONVENIOS","convenio"),("PRIORIDADES","prioridade"),("SIMNAO","aih/autorizada/apa/opme"),("CIRURGIOES","cirurgiao")]
col_of={}
for ci,(key,tit) in enumerate(ordem,1):
    col=get_column_letter(ci); col_of[key]=col
    hc=wv.cell(1,ci,tit); hc.font=HF; hc.fill=BLUE; hc.alignment=Alignment("center")
    for ri,v in enumerate(LISTAS[key],2): wv.cell(ri,ci,v).font=Font(name=F,size=10)
    wv.column_dimensions[col].width=max(14,min(42,max([len(tit)]+[len(str(x)) for x in LISTAS[key]]+[10])+2))

for ci,(h,ob,ex,vk) in enumerate(COLS,1):
    if not vk: continue
    col=col_of[vk]; n=len(LISTAS[vk])
    if n==0: continue
    dv=DataValidation(type="list",formula1=f"='Valores válidos'!${col}$2:${col}${n+1}",allow_blank=True)
    dv.error="Selecione um valor da lista (aba Valores válidos)."; dv.errorTitle="Fora da lista"
    L=get_column_letter(ci); dv.add(f"{L}2:{L}1001"); ws.add_data_validation(dv)

ins=wb.create_sheet("Instruções",0)
linhas=[("MODELO DE IMPORTAÇÃO — FILA CIRÚRGICA",0),
("",1),
("• Preencha 1 linha por cirurgia na aba 'Dados'. A linha 2 é EXEMPLO — apague antes de importar.",1),
("• Cabeçalho VERMELHO = obrigatório. LARANJA = informe CPF OU CNS (ao menos um). AZUL = opcional.",1),
("• Não renomeie/reordene/remova colunas. Não edite a aba 'Valores válidos'.",1),
("",1),
("DATAS: formato AAAA-MM-DD (ex.: 2026-06-20). HORÁRIO: HH:MM. DURACAO: minutos (ex.: 60).",1),
("CPF / CNS / telefone / duração: digite só números; a célula é texto para preservar zeros.",1),
("FLAGS aih/autorizada/apa/opme: SIM ou NAO (o sistema grava como verdadeiro/falso).",1),
("PROCEDIMENTO: digite o nome ou o código SIGTAP; o sistema reconcilia com a tabela oficial.",1),
("CIRURGIÃO / STATUS / etc.: use o menu suspenso. Valor parecido será sugerido; valor novo pede confirmação.",1),
("",1),
("IMPORTAÇÃO EM 2 FASES: 1) o sistema valida tudo e mostra o resultado linha a linha (nada é gravado);",1),
("2) você revisa, resolve as sugestões e confirma. Só então grava.",1),
]
for ri,(t,_) in enumerate(linhas,1):
    c=ins.cell(ri,1,t); c.font=Font(name=F,bold=(ri==1 or t.isupper()),size=14 if ri==1 else (11 if t.endswith(':') else 10), color=("1F4E78" if ri==1 else "000000"))
ins.column_dimensions["A"].width=120

out=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "modelo_importacao_fila_cirurgica.xlsx")
wb.save(out); print("OK",out,"| cirurgioes:",len(CIRURGIOES))
